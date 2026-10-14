"""Calcolare una giornata: dai voti al risultato, e da dove arrivano i voti.

Due cose distinte, tenute separate di proposito:

- **leggere i voti**, che oggi entrano da un file o da un copia-incolla, e un
  domani potrebbero arrivare da una fonte in rete senza che cambi altro;
- **calcolare la giornata**, che prende le formazioni salvate, i voti e le
  regole della lega, e scrive gol e punti nel calendario.

Il calcolo non tocca la rete e non importa Streamlit: si prova per intero coi
test, ed e' giusto cosi', perche' e' il pezzo che decide chi vince.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from .formazioni import (
    SOSTITUZIONI_MASSIME,
    EsitoPartita,
    Formazione,
    Voto,
    calcola_partita,
)
from .leghe import Bonus, FasciaModificatore, ModalitaSostituzioni
from .regole import ParametriLega


class VotiNonLeggibili(ValueError):
    """Il file dei voti non si lascia leggere, e il messaggio dice perche'."""


# Le colonne accettate, per sinonimi: chi prepara il file non deve indovinare
# la nostra grafia. `voto` vuoto vuol dire senza voto, che non e' zero.
_SINONIMI_VOTI = {
    "giocatore": ("giocatore", "nome", "calciatore", "player", "nomegiocatore"),
    "id": ("id", "idgiocatore", "idufficiale", "playerid"),
    "voto": ("voto", "votobase", "rating", "v"),
    "gol": ("gol", "golfatti", "reti", "goals", "gf"),
    "gol_su_rigore": ("golsurigore", "rigorisegnati", "golrigore"),
    "rigori_sbagliati": ("rigorisbagliati", "rigorifalliti", "errori"),
    "rigori_parati": ("rigoriparati", "parate"),
    "autogol": ("autogol", "autoreti", "og"),
    "assist": ("assist", "assistenze", "ass"),
    "ammonizioni": ("ammonizioni", "gialli", "ammonizione", "cartellinigialli"),
    "espulsioni": ("espulsioni", "rossi", "espulsione", "cartellinirossi"),
    "gol_subiti": ("golsubiti", "gs", "retisubite"),
    "imbattuto": ("imbattuto", "cleansheet", "porta inviolata", "portainviolata"),
}

MODELLO_CSV_VOTI = (
    "giocatore;voto;gol;assist;ammonizioni;espulsioni;gol_subiti\n"
    "Svilar;6.5;0;0;0;0;1\n"
    "Dybala;7;1;1;1;0;0\n"
    "Barella;;0;0;0;0;0\n"
)


def _chiave(testo: str) -> str:
    import re
    import unicodedata

    senza_accenti = "".join(
        c
        for c in unicodedata.normalize("NFD", str(testo))
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"[^a-z0-9]", "", senza_accenti.lower())


def _mappa_colonne(intestazioni) -> dict[str, str]:
    normalizzate = {_chiave(c): c for c in intestazioni if c}
    trovate: dict[str, str] = {}
    for nostro, sinonimi in _SINONIMI_VOTI.items():
        for sinonimo in sinonimi:
            if _chiave(sinonimo) in normalizzate:
                trovate[nostro] = normalizzate[_chiave(sinonimo)]
                break
    return trovate


def _intero(valore) -> int:
    testo = str(valore or "").strip().replace(",", ".")
    if not testo:
        return 0
    try:
        return int(float(testo))
    except ValueError:
        return 0


def _booleano(valore) -> bool:
    return _chiave(valore) in ("1", "si", "s", "true", "vero", "x", "y", "yes")


@dataclass
class EsitoLetturaVoti:
    """Cosa e' stato letto e cosa non si e' saputo abbinare."""

    voti: list[Voto] = field(default_factory=list)
    non_abbinati: list[str] = field(default_factory=list)
    senza_voto: int = 0

    @property
    def quanti(self) -> int:
        return len(self.voti)


def leggi_voti(
    contenuto: bytes | str,
    giornata: int,
    per_nome: dict[str, int],
    per_id_ufficiale: dict[int, int] | None = None,
) -> EsitoLetturaVoti:
    """Legge i voti di una giornata da un CSV o da una tabella incollata.

    `per_nome` va da nome normalizzato a id interno del giocatore; se il file
    porta l'id ufficiale del listone si usa quello, che e' piu' sicuro di un
    nome. Chi non si abbina finisce in `non_abbinati` invece di sparire: un
    voto perso in silenzio e' un punteggio sbagliato che nessuno spiega.
    """
    import csv
    import io

    from .importazione import normalizza_nome_giocatore

    testo = (
        contenuto.decode("utf-8-sig", "replace")
        if isinstance(contenuto, bytes)
        else contenuto
    )
    if not testo.strip():
        raise VotiNonLeggibili("Il file dei voti e' vuoto.")

    prima = testo.splitlines()[0]
    conteggi = {c: prima.count(c) for c in ("\t", ";", ",")}
    separatore = max(conteggi, key=lambda c: conteggi[c])
    if not conteggi[separatore]:
        separatore = ";"

    righe = list(csv.DictReader(io.StringIO(testo), delimiter=separatore))
    if not righe:
        raise VotiNonLeggibili(
            "Il file dei voti non ha righe sotto l'intestazione — o non e' un "
            "CSV. Servono almeno le colonne «giocatore» e «voto»."
        )

    colonne = _mappa_colonne(righe[0].keys())
    if "voto" not in colonne or not ({"giocatore", "id"} & set(colonne)):
        lette = ", ".join(c for c in righe[0] if isinstance(c, str) and c.strip())
        raise VotiNonLeggibili(
            "Nel file dei voti servono la colonna del voto e una che dica di "
            f"chi e' (giocatore o id). Intestazioni lette: {lette or '(nessuna)'}."
        )

    esito = EsitoLetturaVoti()
    per_id_ufficiale = per_id_ufficiale or {}
    for riga in righe:

        def campo(nome: str, riga=riga) -> str:
            colonna = colonne.get(nome)
            return str(riga.get(colonna, "") or "").strip() if colonna else ""

        nome = campo("giocatore")
        identificativo = None
        if "id" in colonne and campo("id"):
            identificativo = per_id_ufficiale.get(_intero(campo("id")))
        if identificativo is None and nome:
            identificativo = per_nome.get(normalizza_nome_giocatore(nome))
        if identificativo is None:
            if nome or campo("id"):
                esito.non_abbinati.append(nome or campo("id"))
            continue

        grezzo = campo("voto").replace(",", ".")
        voto = None
        if grezzo:
            try:
                voto = float(grezzo)
            except ValueError:
                voto = None
        if voto is None:
            esito.senza_voto += 1

        esito.voti.append(
            Voto(
                giocatore_id=identificativo,
                giornata=giornata,
                voto=voto,
                gol=_intero(campo("gol")),
                gol_su_rigore=_intero(campo("gol_su_rigore")),
                rigori_sbagliati=_intero(campo("rigori_sbagliati")),
                rigori_parati=_intero(campo("rigori_parati")),
                autogol=_intero(campo("autogol")),
                assist=_intero(campo("assist")),
                ammonizioni=_intero(campo("ammonizioni")),
                espulsioni=_intero(campo("espulsioni")),
                gol_subiti=_intero(campo("gol_subiti")),
                imbattuto=_booleano(campo("imbattuto")),
            )
        )

    if not esito.voti:
        raise VotiNonLeggibili(
            "Nessuna riga abbinata a un giocatore in rosa. Controlla i nomi, "
            "oppure metti nel file la colonna «id» del listone."
        )
    return esito


# --- il calcolo -------------------------------------------------------------


@dataclass
class RisultatoPartita:
    """Una partita calcolata, pronta da scrivere e da mostrare."""

    partita_id: int
    casa_id: int
    trasferta_id: int
    esito: EsitoPartita
    saltata: str = ""

    @property
    def calcolata(self) -> bool:
        return not self.saltata


@dataclass
class EsitoGiornata:
    """Il conto di una giornata intera."""

    giornata: int
    risultati: list[RisultatoPartita] = field(default_factory=list)
    avvisi: list[str] = field(default_factory=list)

    @property
    def calcolate(self) -> int:
        return sum(1 for r in self.risultati if r.calcolata)


def calcola_giornata(
    giornata: int,
    partite: list[dict],
    formazioni: dict[int, Formazione],
    voti: dict[int, Voto],
    ruoli_per_giocatore: dict[int, tuple[str, ...]],
    parametri: ParametriLega,
    bonus: Bonus | None = None,
    fasce_difesa: tuple[FasciaModificatore, ...] = (),
    sostituzioni_massime: int = SOSTITUZIONI_MASSIME,
    modalita: ModalitaSostituzioni = ModalitaSostituzioni.BASIC,
) -> EsitoGiornata:
    """Calcola tutte le partite di una giornata, saltando quelle che non si puo'.

    Una squadra senza formazione non fa perdere l'intera giornata: si salta
    quella partita e lo si dice. Il presidente sistema e ricalcola, invece di
    trovarsi davanti un errore unico che non dice quale squadra manca.
    """
    esito = EsitoGiornata(giornata=giornata)
    if not voti:
        esito.avvisi.append(
            "Non ci sono voti per questa giornata: prima si caricano i voti, "
            "poi si calcola."
        )
        return esito

    for partita in partite:
        casa_id = int(partita["casa_id"])
        trasferta_id = int(partita["trasferta_id"])
        mancanti = [s for s in (casa_id, trasferta_id) if s not in formazioni]
        if mancanti:
            esito.risultati.append(
                RisultatoPartita(
                    partita_id=int(partita["id"]),
                    casa_id=casa_id,
                    trasferta_id=trasferta_id,
                    esito=None,  # type: ignore[arg-type]
                    saltata=(
                        f"{len(mancanti)} squadr"
                        + ("a" if len(mancanti) == 1 else "e")
                        + " senza formazione salvata"
                    ),
                )
            )
            continue

        esito.risultati.append(
            RisultatoPartita(
                partita_id=int(partita["id"]),
                casa_id=casa_id,
                trasferta_id=trasferta_id,
                esito=calcola_partita(
                    formazioni[casa_id],
                    formazioni[trasferta_id],
                    voti,
                    ruoli_per_giocatore,
                    parametri,
                    bonus,
                    fasce_difesa,
                    sostituzioni_massime,
                    modalita,
                ),
            )
        )
    return esito


def abbinamento(giocatori) -> tuple[dict[str, int], dict[int, int]]:
    """Le due chiavi per riconoscere un giocatore in un file di voti.

    La prima va dal nome normalizzato all'id interno, la seconda dall'id del
    listone ufficiale allo stesso id. L'id e' piu' sicuro del nome, ma non
    tutti i file dei voti ce l'hanno, quindi servono tutte e due.
    """
    from .importazione import normalizza_nome_giocatore

    valori = giocatori.values() if hasattr(giocatori, "values") else giocatori
    valori = list(valori)
    return (
        {normalizza_nome_giocatore(g.nome): g.id for g in valori},
        {g.id_ufficiale: g.id for g in valori if g.id_ufficiale},
    )


# L'export «formazioni» di Leghe Fantacalcio
# ---------------------------------------------------------------------------
#
# E' la giornata **gia' giocata e gia' calcolata** dalla piattaforma: per ogni
# partita le due formazioni con modulo, titolari e panchina, il voto e il
# fantavoto di ognuno, i bonus in punti, il modificatore di difesa, il totale
# e il risultato. Non e' un file di voti: e' il tabellino ufficiale.
#
# Due cose stanno **nel colore delle celle** e non in una colonna, e chi legge
# solo i valori le perde:
#
# - la cella dopo il nome, gialla o rossa, e' l'ammonizione o l'espulsione.
#   Nessuna colonna le riporta: il fantavoto le contiene gia', ma ricalcolarlo
#   dalle colonne darebbe mezzo punto in piu' a ogni ammonito;
# - il nome in grigio e' di chi **non e' entrato nel punteggio** — i panchinari
#   rimasti fuori. E' l'unico modo per sapere chi e' entrato davvero al posto
#   di chi non ha preso voto.

GIALLO_AMMONIZIONE = "FFFACC15"
ROSSO_ESPULSIONE = "FFDC2626"
GRIGIO_NON_ENTRATO = "00A0A7B1"

# Le colonne dei bonus, nell'ordine in cui le scrive il file. I valori sono
# **punti**, non conteggi: «G 3» vuol dire tre punti di gol, cioe' un gol con
# la tabella di questa lega.
COLONNE_BONUS = ("G", "GS", "AU", "RS", "RSb", "RP", "A", "AS", "AG")

# Quanto e' larga mezza pagina: la partita sta su due colonne di blocchi, casa
# a sinistra e trasferta a destra.
SCOSTAMENTO_TRASFERTA = 16


@dataclass(frozen=True)
class RigaFormazione:
    """Un giocatore schierato, come lo scrive la piattaforma."""

    ruolo: str
    ruoli: tuple[str, ...]
    nome: str
    club: str
    # Il voto puro, `None` quando e' «SV». Il fantavoto e' quello ufficiale,
    # bonus, malus e cartellini gia' compresi.
    voto: float | None
    fantavoto: float | None
    bonus: dict[str, float]
    ammonito: bool
    espulso: bool
    # Falso per chi e' rimasto in panchina senza entrare nel punteggio.
    conta: bool

    @property
    def ha_giocato(self) -> bool:
        return self.voto is not None


@dataclass
class SquadraInCampo:
    nome: str
    modulo: str
    totale: float
    modificatore_difesa: float
    titolari: list[RigaFormazione] = field(default_factory=list)
    panchina: list[RigaFormazione] = field(default_factory=list)

    @property
    def entrati(self) -> list[RigaFormazione]:
        """Chi dalla panchina e' entrato nel punteggio."""
        return [g for g in self.panchina if g.conta and g.ha_giocato]


@dataclass
class PartitaGiocata:
    casa: SquadraInCampo
    trasferta: SquadraInCampo
    gol_casa: int | None
    gol_trasferta: int | None


@dataclass
class EsitoFormazioniLeghe:
    giornata: int = 0
    partite: list[PartitaGiocata] = field(default_factory=list)
    problemi: list[str] = field(default_factory=list)

    @property
    def squadre(self) -> list[SquadraInCampo]:
        return [s for p in self.partite for s in (p.casa, p.trasferta)]

    @property
    def leggibile(self) -> bool:
        return bool(self.partite) and not self.problemi


def _valore(testo) -> float | None:
    """Un numero del tabellino. «SV» e «-» non sono zero: sono «non c'e'»."""
    if testo is None or str(testo).strip() in ("", "-", "SV"):
        return None
    try:
        return float(str(testo).replace(",", "."))
    except ValueError:
        return None


def _giornata_dal_titolo(testo: str) -> int:
    import re

    trovata = re.search(r"giornata\s+(\d+)", str(testo or ""), re.IGNORECASE)
    return int(trovata.group(1)) if trovata else 0


def _riga_formazione(foglio, riga: int, colonna: int) -> RigaFormazione | None:
    """Legge un giocatore, colore delle celle compreso."""

    def cella(scostamento: int):
        return foglio.cell(row=riga, column=colonna + scostamento)

    ruolo = str(cella(1).value or "").strip()
    if not ruolo or ruolo in ("Panchina", "Modificatori", "Modulo", "Ruolo"):
        return None

    cartellino = cella(3)
    colore_cartellino = (
        str(cartellino.fill.fgColor.rgb or "") if cartellino.fill.patternType else ""
    )
    nome_cella = cella(2)
    colore_nome = (
        str(nome_cella.font.color.rgb or "")
        if nome_cella.font and nome_cella.font.color
        else ""
    )

    return RigaFormazione(
        ruolo=ruolo,
        ruoli=tuple(r.strip() for r in ruolo.split("/") if r.strip()),
        nome=str(nome_cella.value or "").strip(),
        club=str(cella(4).value or "").strip(),
        voto=_valore(cella(5).value),
        fantavoto=_valore(cella(6).value),
        bonus={
            chiave: _valore(cella(7 + n).value) or 0.0
            for n, chiave in enumerate(COLONNE_BONUS)
        },
        ammonito=colore_cartellino == GIALLO_AMMONIZIONE,
        espulso=colore_cartellino == ROSSO_ESPULSIONE,
        conta=colore_nome != GRIGIO_NON_ENTRATO,
    )


def _squadra_in_campo(foglio, riga_risultato: int, colonna: int) -> SquadraInCampo:
    """Un blocco: dal nome della squadra fino al modificatore."""
    squadra = SquadraInCampo(
        nome=str(foglio.cell(row=riga_risultato, column=colonna + 1).value or "").strip(),
        modulo=str(foglio.cell(row=riga_risultato + 2, column=colonna + 2).value or ""),
        totale=_valore(foglio.cell(row=riga_risultato + 2, column=colonna + 6).value)
        or 0.0,
        modificatore_difesa=0.0,
    )

    dove = squadra.titolari
    riga = riga_risultato + 5
    while riga <= foglio.max_row:
        prima = str(foglio.cell(row=riga, column=colonna + 1).value or "").strip()
        if prima == "Panchina":
            dove = squadra.panchina
        elif prima == "Modificatori":
            dove = None
        elif prima == "Modificatore difesa":
            # Quando vale zero la riga non c'e' proprio: e' il motivo per cui
            # il valore parte da zero invece che da «non letto».
            squadra.modificatore_difesa = (
                _valore(foglio.cell(row=riga, column=colonna + 6).value) or 0.0
            )
        elif not prima:
            break
        elif dove is not None:
            letta = _riga_formazione(foglio, riga, colonna)
            if letta is not None:
                dove.append(letta)
        riga += 1
    return squadra


def leggi_formazioni_leghe(contenuto: bytes) -> EsitoFormazioniLeghe:
    """Legge il tabellino di giornata esportato da Leghe Fantacalcio.

    Restituisce le partite com'erano: chi era in campo, chi e' entrato, i voti
    e i totali **della piattaforma**. Qui non si calcola niente — e' il
    tabellino ufficiale, e serve sia per registrarlo sia per confrontarlo con
    il conto che fa il sito.
    """
    import io

    from openpyxl import load_workbook

    esito = EsitoFormazioniLeghe()
    try:
        # Niente `read_only`: ammonizioni e panchinari entrati stanno nel
        # colore delle celle, e in sola lettura openpyxl non lo porta.
        cartella = load_workbook(io.BytesIO(contenuto), data_only=True)
    except Exception as errore:  # noqa: BLE001 - openpyxl alza tipi eterogenei
        esito.problemi.append(f"Non e' un file Excel leggibile: {errore}")
        return esito

    foglio = cartella[cartella.sheetnames[0]]
    esito.giornata = _giornata_dal_titolo(
        foglio.cell(row=1, column=1).value
    ) or _giornata_dal_titolo(foglio.title)

    for riga in range(1, foglio.max_row + 1):
        risultato = str(foglio.cell(row=riga, column=16).value or "")
        casa = foglio.cell(row=riga, column=1).value
        if not casa or "-" not in risultato:
            continue
        gol = [_valore(p) for p in risultato.split("-")]
        esito.partite.append(
            PartitaGiocata(
                casa=_squadra_in_campo(foglio, riga, 0),
                trasferta=_squadra_in_campo(foglio, riga, SCOSTAMENTO_TRASFERTA),
                gol_casa=int(gol[0]) if len(gol) == 2 and gol[0] is not None else None,
                gol_trasferta=(
                    int(gol[1]) if len(gol) == 2 and gol[1] is not None else None
                ),
            )
        )

    if not esito.partite:
        esito.problemi.append(
            "Nessuna partita leggibile: serve il file «formazioni» esportato "
            "da Leghe Fantacalcio, quello con il tabellino di ogni squadra."
        )
    return esito


def _conteggio(punti: float, valore_unitario: float) -> int:
    """Quanti eventi valgono quei punti, con la tabella di questa lega.

    Il file scrive «G 3» e non «1 gol»: per tornare al numero di gol serve il
    valore che la lega da' al gol. Se quel bonus vale zero il conto non si
    puo' fare, e si preferisce zero eventi a una divisione per zero.
    """
    if not punti or not valore_unitario:
        return 0
    return int(round(punti / valore_unitario))


def voto_dal_tabellino(
    riga: RigaFormazione,
    giocatore_id: int,
    giornata: int,
    bonus: Bonus | None = None,
) -> Voto:
    """Dal tabellino della piattaforma al `Voto` che il dominio sa calcolare.

    Serve a confrontare il nostro conto con il loro: ricostruisce gli eventi
    (gol, assist, cartellini) dai punti scritti nel file, cosi' il punteggio
    lo rifa' la lega con la **sua** tabella invece di copiare il fantavoto.
    Quando interessa solo il risultato ufficiale, `riga.fantavoto` lo dice
    gia' e non serve passare di qui.
    """
    bonus = bonus or Bonus()
    punti = riga.bonus
    return Voto(
        giocatore_id=giocatore_id,
        giornata=giornata,
        voto=riga.voto,
        gol=_conteggio(punti.get("G", 0), bonus.gol_segnato),
        gol_su_rigore=_conteggio(punti.get("RS", 0), bonus.gol_su_rigore),
        rigori_sbagliati=_conteggio(punti.get("RSb", 0), bonus.rigore_sbagliato),
        rigori_parati=_conteggio(punti.get("RP", 0), bonus.rigore_parato),
        autogol=_conteggio(punti.get("AU", 0), bonus.autogol),
        # Il file distingue assist, assist soft e assist gold; la tabella
        # della lega ha un valore solo, quindi i tre si sommano in punti e si
        # riconvertono una volta sola.
        assist=_conteggio(
            punti.get("A", 0) + punti.get("AS", 0) + punti.get("AG", 0), bonus.assist
        ),
        ammonizioni=1 if riga.ammonito else 0,
        espulsioni=1 if riga.espulso else 0,
        gol_subiti=_conteggio(punti.get("GS", 0), bonus.gol_subito),
    )
